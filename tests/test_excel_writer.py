from __future__ import annotations

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from autopo.config import DATE_COLUMNS
from autopo.core.excel_writer import append_rows, ensure_workbook

ROW = {
    "Upload": "Y",
    "Sold-to Party": "10001",
    "Customer Reference": "242445",
    "Customer Material": "C460-3373",
    "Order Quantity": "40",
    "Unit Price": "1234.56",
    "CRD": "2025/04/07",
    "Customer Reference Date": "2025/04/19",
}


@pytest.fixture()
def workbook_path(tmp_path: Path) -> Path:
    return tmp_path / "open_order.xlsx"


def headers_of(path: Path, sheet: str = "OpenOrder") -> list[str]:
    ws = load_workbook(path)[sheet]
    return [c.value for c in ws[1] if c.value]


class TestEnsureWorkbook:
    def test_creates_a_sheet_with_headers(self, workbook_path: Path):
        wb = ensure_workbook(workbook_path)
        ws = wb["OpenOrder"]
        assert ws["A1"].value == "Upload"
        assert ws.freeze_panes == "A2"

    def test_reuses_an_existing_workbook(self, workbook_path: Path):
        append_rows(workbook_path, [ROW])
        wb = ensure_workbook(workbook_path)
        assert wb["OpenOrder"].max_row == 2

    def test_adds_a_missing_sheet_to_an_existing_workbook(self, workbook_path: Path):
        append_rows(workbook_path, [ROW])
        wb = ensure_workbook(workbook_path, sheet_name="EU")
        assert "EU" in wb.sheetnames


class TestAppendRows:
    def test_creates_the_output_directory(self, tmp_path: Path):
        path = tmp_path / "nested" / "dir" / "open_order.xlsx"
        append_rows(path, [ROW])
        assert path.exists()

    def test_returns_the_number_of_rows_written(self, workbook_path: Path):
        assert append_rows(workbook_path, [ROW, ROW, ROW]) == 3

    def test_writes_values_under_the_matching_header(self, workbook_path: Path):
        append_rows(workbook_path, [ROW])
        ws = load_workbook(workbook_path)["OpenOrder"]
        by_header = dict(zip(headers_of(workbook_path), [c.value for c in ws[2]]))
        assert by_header["Sold-to Party"] == "10001"
        assert by_header["Order Quantity"] == "40"
        assert by_header["CRD"] == "2025/04/07"

    def test_columns_the_row_does_not_fill_are_left_blank(self, workbook_path: Path):
        append_rows(workbook_path, [ROW])
        ws = load_workbook(workbook_path)["OpenOrder"]
        by_header = dict(zip(headers_of(workbook_path), [c.value for c in ws[2]]))
        assert by_header["Remark"] is None

    def test_appending_twice_accumulates_instead_of_overwriting(self, workbook_path: Path):
        # A day's run appends several POs to the same tracker.
        append_rows(workbook_path, [ROW, ROW])
        append_rows(workbook_path, [ROW])
        ws = load_workbook(workbook_path)["OpenOrder"]
        assert ws.max_row == 4  # 1 header + 3 data rows

    def test_the_header_row_is_written_once(self, workbook_path: Path):
        append_rows(workbook_path, [ROW])
        append_rows(workbook_path, [ROW])
        ws = load_workbook(workbook_path)["OpenOrder"]
        assert [c.value for c in ws[2]] != [c.value for c in ws[1]]
        assert ws["A1"].value == "Upload"

    def test_writes_to_the_requested_sheet(self, workbook_path: Path):
        append_rows(workbook_path, [ROW], sheet_name="EU")
        assert "EU" in load_workbook(workbook_path).sheetnames

    def test_a_second_sheet_gets_its_own_headers(self, workbook_path: Path):
        append_rows(workbook_path, [ROW])
        append_rows(workbook_path, [ROW], sheet_name="EU")
        assert headers_of(workbook_path, "EU") == headers_of(workbook_path, "OpenOrder")

    def test_writing_no_rows_leaves_a_headers_only_workbook(self, workbook_path: Path):
        assert append_rows(workbook_path, []) == 0
        assert load_workbook(workbook_path)["OpenOrder"].max_row == 1


class TestErpImportFormat:
    """The importer this workbook feeds rejects typed cells.

    Nothing may be coerced to a number or a date on the way in -- not a
    quantity, not a price, not a CRD -- or the upload is refused. These tests
    exist so that "fixing" the types, which looks like an obvious improvement
    from inside Excel, fails loudly here instead of silently at upload.
    """

    def cells_of(self, path: Path, header: str):
        ws = load_workbook(path)["OpenOrder"]
        column = headers_of(path).index(header) + 1
        return [ws.cell(row=r, column=column) for r in range(2, ws.max_row + 1)]

    def test_every_populated_cell_is_text(self, workbook_path: Path):
        append_rows(workbook_path, [ROW])
        ws = load_workbook(workbook_path)["OpenOrder"]
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.value is not None:
                    assert cell.data_type == "s", f"{cell.coordinate} is not text"

    def test_a_quantity_stays_text_although_it_looks_numeric(self, workbook_path: Path):
        append_rows(workbook_path, [ROW])
        (qty,) = self.cells_of(workbook_path, "Order Quantity")
        assert qty.value == "40"
        assert qty.data_type == "s"

    def test_a_price_stays_text(self, workbook_path: Path):
        append_rows(workbook_path, [ROW])
        (price,) = self.cells_of(workbook_path, "Unit Price")
        assert price.value == "1234.56"
        assert price.data_type == "s"

    @pytest.mark.parametrize("header", ["CRD", "Customer Reference Date"])
    def test_a_date_column_is_text_in_yyyy_mm_dd_form(
        self, workbook_path: Path, header: str
    ):
        append_rows(workbook_path, [ROW])
        (cell,) = self.cells_of(workbook_path, header)
        assert cell.data_type == "s"
        assert re.fullmatch(r"\d{4}/\d{2}/\d{2}", cell.value)

    def test_no_number_format_is_applied(self, workbook_path: Path):
        # A number format on a text cell does nothing; one here would only
        # mislead the next reader into thinking these are real dates.
        append_rows(workbook_path, [ROW])
        (crd,) = self.cells_of(workbook_path, "CRD")
        assert crd.number_format == "General"

    def test_a_value_that_is_not_a_string_is_coerced(self, workbook_path: Path):
        # Nothing produces these today. The guarantee is enforced at the
        # boundary so that it survives whatever a future parser returns.
        append_rows(workbook_path, [{**ROW, "Order Quantity": 40}])
        (qty,) = self.cells_of(workbook_path, "Order Quantity")
        assert qty.value == "40"
        assert qty.data_type == "s"

    def test_every_date_column_is_a_real_header(self, workbook_path: Path):
        # Guards against a rename in config.py drifting from the writer.
        append_rows(workbook_path, [ROW])
        assert set(DATE_COLUMNS) - {"ETA"} <= set(headers_of(workbook_path))
