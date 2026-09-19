"""The loop the CLI and the GUI share.

These tests used to be reachable only through the CLI; the GUI carried its own
copy and was covered by nothing.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from autopo.core.pipeline import FileResult, IngestSummary, collect_pdfs, ingest
from tests.factories import (
    ItemA,
    build_a_pdf,
    build_b_pdf,
    build_plain_pdf,
    build_unreadable_pdf,
)

B_ROWS = [
    ["CB-64810", "2025/04/19", "C948-2452", "SKU-7919-E85", "59", "450", "USD", "2025/06/02"],
]


@pytest.fixture()
def inbox(tmp_path: Path) -> Path:
    folder = tmp_path / "inbox"
    folder.mkdir()
    build_a_pdf(
        folder / "customer_a.pdf",
        [
            ItemA("C460-3373", "40", "1,234.56", "49,382.40", "Apr 19, 2025"),
            ItemA("C650-7790", "7", "99.00", "693.00", "May 02, 2025"),
        ],
    )
    build_b_pdf(folder / "customer_b.pdf", B_ROWS)
    return folder


@pytest.fixture()
def workbook(tmp_path: Path) -> Path:
    return tmp_path / "out" / "open_order.xlsx"


class TestCollectPdfs:
    def test_lists_a_directory_in_a_stable_order(self, inbox: Path):
        assert [p.name for p in collect_pdfs(inbox)] == [
            "customer_a.pdf",
            "customer_b.pdf",
        ]

    def test_a_single_file_collects_to_itself(self, inbox: Path):
        pdf = inbox / "customer_b.pdf"
        assert collect_pdfs(pdf) == [pdf]

    def test_non_pdfs_are_ignored(self, inbox: Path):
        (inbox / "notes.txt").write_text("nothing to see")
        assert len(collect_pdfs(inbox)) == 2

    def test_the_pdf_extension_is_matched_case_insensitively(self, tmp_path: Path):
        folder = tmp_path / "shouty"
        folder.mkdir()
        build_b_pdf(folder / "CUSTOMER_B.PDF", B_ROWS)
        assert len(collect_pdfs(folder)) == 1

    def test_an_empty_folder_collects_nothing(self, tmp_path: Path):
        empty = tmp_path / "empty"
        empty.mkdir()
        assert collect_pdfs(empty) == []

    def test_a_missing_path_raises(self, tmp_path: Path):
        # Each front-end renders this in its own idiom, so the pipeline raises.
        with pytest.raises(FileNotFoundError):
            collect_pdfs(tmp_path / "nope")


class TestIngest:
    def test_writes_every_line_item(self, inbox: Path, workbook: Path):
        summary = ingest(collect_pdfs(inbox), workbook)
        assert summary.total_rows == 3
        assert load_workbook(workbook)["OpenOrder"].max_row == 4

    def test_reports_the_customer_and_the_sku_matches_per_file(
        self, inbox: Path, workbook: Path
    ):
        summary = ingest(collect_pdfs(inbox), workbook)
        labels = {r.customer_label for r in summary.results}
        assert labels == {"Customer-A", "Customer-B"}
        assert all(r.matched == r.rows for r in summary.results)

    def test_calls_back_once_per_file_as_it_goes(self, inbox: Path, workbook: Path):
        # The GUI relies on this to show progress during a long run.
        seen: list[FileResult] = []
        ingest(collect_pdfs(inbox), workbook, on_file=seen.append)
        assert [r.path.name for r in seen] == ["customer_a.pdf", "customer_b.pdf"]

    def test_the_callback_is_optional(self, inbox: Path, workbook: Path):
        assert ingest(collect_pdfs(inbox), workbook).total_rows == 3

    def test_writes_to_the_requested_sheet(self, inbox: Path, workbook: Path):
        ingest(collect_pdfs(inbox), workbook, sheet="EU")
        assert "EU" in load_workbook(workbook).sheetnames

    def test_an_unrecognised_pdf_is_skipped_with_a_reason(
        self, inbox: Path, workbook: Path
    ):
        build_plain_pdf(inbox / "memo.pdf")
        summary = ingest(collect_pdfs(inbox), workbook)

        (skipped,) = summary.skipped
        assert skipped.path.name == "memo.pdf"
        assert "No parser fingerprint matched" in skipped.skipped
        assert skipped.rows == 0

    def test_one_skipped_file_does_not_cost_the_others(self, inbox: Path, workbook: Path):
        build_plain_pdf(inbox / "memo.pdf")
        assert ingest(collect_pdfs(inbox), workbook).total_rows == 3

    def test_a_damaged_file_is_skipped_rather_than_aborting_the_batch(
        self, inbox: Path, workbook: Path
    ):
        # A truncated or encrypted PDF used to raise straight out of the loop.
        # Since the files are processed in name order, one bad file early in
        # the alphabet meant no workbook at all.
        build_unreadable_pdf(inbox / "aaa_damaged.pdf")
        summary = ingest(collect_pdfs(inbox), workbook)

        assert summary.total_rows == 3
        (skipped,) = summary.skipped
        assert skipped.path.name == "aaa_damaged.pdf"

    def test_the_skip_reason_names_the_underlying_error(
        self, inbox: Path, workbook: Path
    ):
        # The operator has to be able to tell "we have no parser for this" from
        # "this file is broken" without reading a traceback.
        build_unreadable_pdf(inbox / "damaged.pdf")
        build_plain_pdf(inbox / "memo.pdf")
        reasons = {r.path.name: r.skipped for r in ingest(collect_pdfs(inbox), workbook).skipped}

        assert "PdfminerException" in reasons["damaged.pdf"]
        assert "No parser fingerprint matched" in reasons["memo.pdf"]

    def test_every_good_file_after_a_damaged_one_is_still_processed(
        self, inbox: Path, workbook: Path
    ):
        build_unreadable_pdf(inbox / "aaa_damaged.pdf")
        results = ingest(collect_pdfs(inbox), workbook).results
        assert [r.rows for r in results] == [0, 2, 1]

    def test_real_parser_output_reaches_the_workbook_as_text(
        self, inbox: Path, workbook: Path
    ):
        # The end-to-end version of the guarantee in core/excel_writer: the ERP
        # importer rejects typed cells, so nothing a parser produces -- through
        # normalisation, SKU enrichment and the writer -- may arrive as a
        # number or a date.
        ingest(collect_pdfs(inbox), workbook)
        ws = load_workbook(workbook)["OpenOrder"]
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.value is not None:
                    assert cell.data_type == "s", f"{cell.coordinate} is not text"

    def test_ingesting_nothing_is_not_an_error(self, workbook: Path):
        summary = ingest([], workbook)
        assert summary.total_rows == 0
        assert not workbook.exists()


class TestSummary:
    def test_totals_are_derived_from_the_results(self):
        summary = IngestSummary(
            [
                FileResult(Path("a.pdf"), "Customer-A", rows=5, matched=5),
                FileResult(Path("b.pdf"), skipped="no fingerprint"),
            ]
        )
        assert summary.total_rows == 5
        assert [r.path.name for r in summary.skipped] == ["b.pdf"]

    def test_an_empty_summary(self):
        assert IngestSummary().total_rows == 0

    def test_a_parsed_file_is_not_marked_skipped(self):
        assert not FileResult(Path("a.pdf"), "Customer-A", rows=5).was_skipped
