"""End-to-end coverage of `python -m autopo.cli ingest ...`."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from autopo.cli import main
from tests.factories import ItemA, build_a_pdf, build_b_pdf, build_plain_pdf

B_ROWS = [
    ["CB-64810", "2025/04/19", "C948-2452", "SKU-7919-E85", "59", "450", "USD", "2025/06/02"],
]


@pytest.fixture()
def inbox(tmp_path: Path) -> Path:
    """A drop folder holding one PO of each format."""
    folder = tmp_path / "inbox"
    folder.mkdir()
    build_a_pdf(
        folder / "customer_a.pdf",
        [
            ItemA("C460-3373", "40", "1,234.56", "49,382.40", "Apr 19, 2025"),
            ItemA("C650-7790", "7", "99.00", "693.00", "May 02, 2025"),
        ],
    )
    build_b_pdf(folder / "customer_b.pdf", B_ROWS)
    return folder


@pytest.fixture()
def workbook(tmp_path: Path) -> Path:
    return tmp_path / "out" / "open_order.xlsx"


def run(*argv: str) -> int:
    return main(list(argv))


class TestIngest:
    def test_exits_zero(self, inbox: Path, workbook: Path):
        assert run("ingest", str(inbox), "--workbook", str(workbook)) == 0

    def test_writes_every_line_item_from_every_pdf(self, inbox: Path, workbook: Path):
        run("ingest", str(inbox), "--workbook", str(workbook))
        ws = load_workbook(workbook)["OpenOrder"]
        assert ws.max_row == 4  # 1 header + 2 Customer-A lines + 1 Customer-B line

    def test_resolves_internal_material_numbers_along_the_way(
        self, inbox: Path, workbook: Path
    ):
        run("ingest", str(inbox), "--workbook", str(workbook))
        ws = load_workbook(workbook)["OpenOrder"]
        headers = [c.value for c in ws[1]]
        materials = {
            ws.cell(row=r, column=headers.index("Material") + 1).value
            for r in range(2, ws.max_row + 1)
        }
        assert "ATP-1042-A17" in materials

    def test_accepts_a_single_pdf(self, inbox: Path, workbook: Path):
        assert run("ingest", str(inbox / "customer_b.pdf"), "--workbook", str(workbook)) == 0
        assert load_workbook(workbook)["OpenOrder"].max_row == 2

    def test_writes_to_the_requested_sheet(self, inbox: Path, workbook: Path):
        run("ingest", str(inbox), "--workbook", str(workbook), "--sheet", "EU")
        assert "EU" in load_workbook(workbook).sheetnames

    def test_reports_each_file_and_the_total(self, inbox: Path, workbook: Path, capsys):
        run("ingest", str(inbox), "--workbook", str(workbook))
        out = capsys.readouterr().out
        assert "customer_a.pdf" in out
        assert "Customer-B" in out
        assert "Wrote 3 row(s)" in out

    def test_a_second_run_appends_rather_than_replacing(self, inbox: Path, workbook: Path):
        run("ingest", str(inbox), "--workbook", str(workbook))
        run("ingest", str(inbox), "--workbook", str(workbook))
        assert load_workbook(workbook)["OpenOrder"].max_row == 7


class TestUnhappyPaths:
    def test_unrecognised_pdf_is_skipped_without_failing_the_batch(
        self, inbox: Path, workbook: Path, capsys
    ):
        # One odd file in the drop folder must not cost the operator the run.
        build_plain_pdf(inbox / "memo.pdf")
        assert run("ingest", str(inbox), "--workbook", str(workbook)) == 0
        assert "[skip] memo.pdf" in capsys.readouterr().out
        assert load_workbook(workbook)["OpenOrder"].max_row == 4

    def test_missing_source_exits_two(self, tmp_path: Path, workbook: Path, capsys):
        assert run("ingest", str(tmp_path / "nope"), "--workbook", str(workbook)) == 2
        assert "does not exist" in capsys.readouterr().err

    def test_folder_without_pdfs_exits_one(self, tmp_path: Path, workbook: Path):
        empty = tmp_path / "empty"
        empty.mkdir()
        (empty / "notes.txt").write_text("nothing to see")
        assert run("ingest", str(empty), "--workbook", str(workbook)) == 1

    def test_no_workbook_is_created_when_there_is_nothing_to_ingest(
        self, tmp_path: Path, workbook: Path
    ):
        empty = tmp_path / "empty"
        empty.mkdir()
        run("ingest", str(empty), "--workbook", str(workbook))
        assert not workbook.exists()

    def test_pdf_extension_matching_is_case_insensitive(self, tmp_path: Path, workbook: Path):
        folder = tmp_path / "shouty"
        folder.mkdir()
        build_b_pdf(folder / "CUSTOMER_B.PDF", B_ROWS)
        assert run("ingest", str(folder), "--workbook", str(workbook)) == 0

    def test_a_subcommand_is_required(self):
        with pytest.raises(SystemExit):
            main([])
